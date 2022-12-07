from plxscripting.easy import *
import math
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import (
    ScatterChart,
    Reference,
    Series,
)
###############################################
PORT_i = 10000  # Define a port number.
PORT_o = 10001
PASSWORD = 'SxDBR<TYKRAX834~'  # Define a password.


# Start the scripting server.
# global g_i, s_i
s_i, g_i = new_server('localhost', PORT_i, password=PASSWORD)
s_o, g_o = new_server('localhost', PORT_o, password=PASSWORD)

###############################################
#Elements:
EXCEL_PATH=r'C:\Users\phtsang\Desktop\PLAXIS_V22\Script\\'
EXCEL_NAME='Plate_y.xlsx'

FILENAME=EXCEL_PATH+EXCEL_NAME

plate=[plt for plt in g_o.Plates[:]]
phase=[p for p in g_o.Phases[:]]

###############################################
#Inputs:
plate_input=['Plate_1']
phase_input=['Installation of strut [Phase_3]','Second (submerged) excavation stage [Phase_4]','Third excavation stage [Phase_5]']

###############################################
#Get Plate results:
def get_plate(plate_o,phase_o):

    plateY=g_o.getresults(plate_o,phase_o,g_o.ResultTypes.Plate.Y, "node")
    plateM=g_o.getresults(plate_o,phase_o,g_o.ResultTypes.Plate.M2D, "node")

    phasename=str(phase_o.Identification).split('[')[0]
    col1='Bending Moment [kNm/m]'+'_'+phasename

    results = {'Y': plateY,col1: plateM}

    plateresults=pd.DataFrame(results)
    plateresults = plateresults.sort_values(by=['Y'],ascending=False)

    return plateresults
###############################################
#Export Excel:
def export_excel(plate_input,phase_input,filename):
    writer = pd.ExcelWriter(filename, engine='xlsxwriter')
    combined=[]
    for i in range(len(phase)):
        for j in range(len(phase_input)):
            if phase[i].Identification == phase_input[j]:
                name=str(phase[i].Identification).split(' [')[1]
                name=name.split(']')[0]
                sheet_name = "%s_%s" % (plate[0].Name, name)
                results = get_plate(plate[0], phase[i])
                combined.append(results)
                results.to_excel(writer,sheet_name=sheet_name,index=False)

    combinedsheet='combined'+'_'+str(plate[0].Name)
    combined=pd.concat(combined, axis=1)
    combined.to_excel(writer,sheet_name=combinedsheet,index=False)
    writer.save()


export_excel(plate_input,phase_input,FILENAME)


def get_combined_plot(filename,sheetname):
    df_inter = pd.read_excel(filename, sheet_name = sheetname,engine="openpyxl")
    wb=load_workbook(filename)
    sheet=wb[sheetname]

    chart1=ScatterChart()
    chart1.x_axis.title = 'Bending Moment (kNm/m)'
    chart1.y_axis.title = 'RL (m)'
    chart={'chart1':chart1} 
    yvalue=Reference(sheet,min_col=1,min_row=2,max_row=len(df_inter)+1)
    position='G'
    prow=1

    if df_inter.columns.values[1].split(' [')[0] == 'Bending Moment':
        value=Reference(sheet,min_col=2,min_row=2,max_row=len(df_inter)+1)
        series=Series(yvalue,value,title=list(df_inter.columns.values)[1])

        chart1.series.append(series)
        
    charts='chart1'
    chart[charts].height=15
    chart[charts].y_axis.tickLblPos = 'low'
    chart[charts].legend.position = 'b'

    if ord(position)<89 and prow<=2:
        sheet.add_chart(chart[charts], position+str(1))
    position=chr(ord(position)+10)
    prow=prow+1
    wb.save(filename)

    for i in range(3,len(df_inter.columns)+1):
        if df_inter.columns.values[i-1].split('.')[0] != 'Y':
            if df_inter.columns.values[i-1].split(' [')[0] == 'Bending Moment':
                value=Reference(sheet,min_col=i,min_row=2,max_row=len(df_inter)+1)
                series=Series(yvalue,value,title=list(df_inter.columns.values)[i-1])
                chart1.series.append(series)
        elif df_inter.columns.values[i-1].split('.')[0] == 'Y':
            yvalue=Reference(sheet,min_col=i,min_row=2,max_row=len(df_inter)+1)      
    wb.save(filename)                    

combinedsheet='combined_Plate_1'
get_combined_plot(FILENAME,combinedsheet)


